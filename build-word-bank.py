import json
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).parent
raw = (ROOT / "presenter-data.js").read_text(encoding="utf-8")
questions = json.loads(re.search(r"=\s*(\[.*\]);\s*$", raw, re.S).group(1))

labels = {
    "easy": "Fácil",
    "medium": "Médio",
    "medium-hard": "Médio-difícil",
    "hard": "Difícil",
}

categories = [
    ("general", "1. Gerais", "Conhecimentos gerais organizados do nível fácil ao difícil."),
    ("sports", "2. Esportes", "Perguntas esportivas organizadas do nível fácil ao difícil."),
    ("portuguese", "3. Língua Portuguesa", "Gramática, literatura e história da língua portuguesa."),
    ("quotes", "4. Quem disse essa frase?", "Frases com quatro alternativas de autoria."),
]

apostles = {
    "question": "Quais são os doze Apóstolos de Nosso Senhor Jesus Cristo?",
    "difficulty": "medium",
    "options": [
        "Simão Pedro, André, Tiago Maior, João, Filipe, Bartolomeu, Mateus, Tomé, Tiago Menor, Judas Tadeu, Simão Zelote e Judas Iscariotes.",
        "Simão Pedro, André, Tiago Maior, João, Filipe, Barnabé, Mateus, Tomé, Tiago Menor, Judas Tadeu, Simão Zelote e São Paulo.",
        "Simão Pedro, André, Tiago Maior, João, Filipe, Bartolomeu, Lucas, Marcos, Tiago Menor, Judas Tadeu, Simão Zelote e Matias.",
        "Simão Pedro, André, Tiago Maior, João, Filipe, Bartolomeu, Mateus, Tomé, Estêvão, Judas Tadeu, Paulo e Matias.",
    ],
    "answer": "Simão Pedro, André, Tiago Maior, João, Filipe, Bartolomeu, Mateus, Tomé, Tiago Menor, Judas Tadeu, Simão Zelote e Judas Iscariotes (posteriormente substituído por São Matias).",
}


def shade(cell, fill):
    props = cell._tc.get_or_add_tcPr()
    element = OxmlElement("w:shd")
    element.set(qn("w:fill"), fill)
    props.append(element)


def set_cell_margins(cell, top=150, start=180, bottom=150, end=180):
    props = cell._tc.get_or_add_tcPr()
    margins = props.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        props.append(margins)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = OxmlElement(f"w:{margin}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        margins.append(node)


def add_question(document, number, item):
    table = document.add_table(rows=1, cols=1)
    table.autofit = True
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    cell = table.cell(0, 0)
    shade(cell, "F4F7F5")
    set_cell_margins(cell)

    heading = cell.paragraphs[0]
    heading.paragraph_format.space_after = Pt(5)
    run = heading.add_run(f"{number}. {item['question']}")
    run.bold = True
    run.font.size = Pt(11)
    badge = heading.add_run(f"   [{labels[item['difficulty']]}]")
    badge.bold = True
    badge.font.size = Pt(8.5)
    badge.font.color.rgb = RGBColor(32, 91, 78)

    for letter, option in zip("ABCD", item["options"]):
        paragraph = cell.add_paragraph()
        paragraph.paragraph_format.left_indent = Cm(0.45)
        paragraph.paragraph_format.space_after = Pt(1.5)
        marker = paragraph.add_run(f"{letter}) ")
        marker.bold = True
        marker.font.color.rgb = RGBColor(32, 91, 78)
        paragraph.add_run(option)

    answer = cell.add_paragraph()
    answer.paragraph_format.space_before = Pt(5)
    answer.paragraph_format.space_after = Pt(0)
    answer_run = answer.add_run("Resposta: ")
    answer_run.bold = True
    answer_run.font.color.rgb = RGBColor(126, 77, 12)
    answer.add_run(item["answer"]).bold = True
    document.add_paragraph().paragraph_format.space_after = Pt(1)


document = Document()
section = document.sections[0]
section.top_margin = Cm(1.7)
section.bottom_margin = Cm(1.7)
section.left_margin = Cm(1.8)
section.right_margin = Cm(1.8)

styles = document.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10)
styles["Title"].font.name = "Georgia"
styles["Title"].font.size = Pt(28)
styles["Title"].font.color.rgb = RGBColor(24, 70, 61)
styles["Heading 1"].font.name = "Georgia"
styles["Heading 1"].font.size = Pt(21)
styles["Heading 1"].font.color.rgb = RGBColor(24, 70, 61)

title = document.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run("Banco de Perguntas")
subtitle = document.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.add_run("Quiz Presencial · Festa de São Pio X").bold = True
summary = document.add_paragraph()
summary.alignment = WD_ALIGN_PARAGRAPH.CENTER
summary.add_run("87 perguntas com alternativas, respostas e dificuldade atribuída.").italic = True
document.add_page_break()

order = {"easy": 0, "medium": 1, "medium-hard": 2, "hard": 3}
for category_index, (topic, title_text, description) in enumerate(categories):
    if category_index:
        document.add_page_break()
    document.add_heading(title_text, level=1)
    document.add_paragraph(description)
    items = sorted(
        (item for item in questions if item["topic"] == topic),
        key=lambda item: (order[item["difficulty"]], int(item["id"].rsplit("-", 1)[1])),
    )
    for index, item in enumerate(items, 1):
        add_question(document, index, item)

document.add_page_break()
document.add_heading("5. Quais são os 12 Apóstolos?", level=1)
document.add_paragraph("Pergunta especial sobre os primeiros discípulos escolhidos por Nosso Senhor.")
add_question(document, 1, apostles)

for section in document.sections:
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Festa de São Pio X · Banco de perguntas")

properties = document.core_properties
properties.title = "Banco de Perguntas — Quiz Presencial"
properties.subject = "Perguntas, alternativas, respostas e níveis de dificuldade"
properties.author = "Associação São José"
document.save(ROOT / "perguntas.docx")
print("perguntas.docx criado com 87 perguntas em 5 categorias.")
